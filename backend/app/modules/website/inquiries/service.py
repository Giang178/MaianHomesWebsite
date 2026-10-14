import csv
import hashlib
import io
import json
import logging
import re
import zipfile
from datetime import datetime, timedelta, timezone

from bson import ObjectId
from fastapi import HTTPException
from pymongo import ReturnDocument
from pymongo.errors import DuplicateKeyError

from app.core.config import settings
from app.db.mongodb import get_database
from app.modules.content.media.service import marketing_media_service
from app.modules.website.common.transactions import session_kwargs, website_transaction

logger = logging.getLogger(__name__)


def utcnow():
    return datetime.now(timezone.utc)


def safe_spreadsheet_value(value):
    """Prevent customer-controlled cells from executing as spreadsheet formulas."""
    if isinstance(value, str):
        starts_with_control = value.startswith(("\t", "\r"))
        starts_with_formula = value.lstrip().startswith(("=", "+", "-", "@"))
        if starts_with_control or starts_with_formula:
            return "'" + value
    return value


FORM_LABELS: dict[str, str] = {
    "COWORKING_CONSULTATION_HERO": "Đăng ký tư vấn Coworking (Hero)",
    "coworking-consultation-modal": "Đăng ký xem chỗ ngồi Coworking",
    "coworking-inquiry": "Tư vấn gói dịch vụ Coworking",
    "office-inquiry": "Tư vấn thuê văn phòng trọn gói",
    "consignment": "Ký gửi Bất động sản",
    "contact_form": "Form liên hệ trực tiếp",
    "website-contact": "Liên hệ chung từ website",
    "project-inquiry": "Đăng ký tư vấn dự án",
    "maya-chat-lead": "Tư vấn qua Chatbot Mái Đỏ",
}

SOURCE_LABELS: dict[str, str] = {
    "COWORKING": "Khách Co-working",
    "OFFICE": "Khách Thuê Văn phòng",
    "PROJECT": "Khách Dự án BĐS",
    "CONSIGNMENT": "Khách Ký gửi BĐS",
    "WEBSITE": "Khách Liên hệ chung",
    "GENERAL": "Chung",
}

STATUS_CONFIG: dict[str, dict[str, str]] = {
    "NEW": {"label": "Chờ tiếp nhận", "color": "#E11D48"},
    "VIEWED": {"label": "Đã tiếp nhận", "color": "#2563EB"},
    "EXPORTED": {"label": "Đã xuất dữ liệu", "color": "#059669"},
    "ARCHIVED": {"label": "Lưu trữ", "color": "#64748B"},
}


def format_price_value(val) -> str | None:
    if val is None or val == "":
        return None
    if isinstance(val, (int, float)):
        if val >= 1_000_000_000:
            ty = val / 1_000_000_000
            return f"{ty:g} Tỷ VNĐ"
        elif val >= 1_000_000:
            tr = val / 1_000_000
            return f"{tr:g} Triệu VNĐ"
        return f"{val:,.0f} VNĐ"
    return str(val).strip()


def format_property_type(val: str | None) -> str | None:
    if not val:
        return None
    pt_map = {
        "apartment": "Căn hộ chung cư",
        "house": "Nhà phố",
        "villa": "Biệt thự",
        "land": "Đất nền",
        "commercial": "Bất động sản thương mại",
    }
    return pt_map.get(val.lower(), val)


def format_transaction_type(val: str | None) -> str | None:
    if not val:
        return None
    tx_map = {
        "sale": "Bán",
        "rent": "Cho thuê",
        "sale_and_rent": "Bán & Cho thuê",
    }
    return tx_map.get(val.lower(), val)


def serialize(document: dict, media_map: dict | None = None) -> dict:
    result = dict(document)
    result["id"] = str(result.pop("_id"))

    if isinstance(result.get("submitted_at"), datetime):
        result["submitted_at"] = result["submitted_at"].isoformat()
    if isinstance(result.get("created_at"), datetime):
        result["created_at"] = result["created_at"].isoformat()
    if isinstance(result.get("updated_at"), datetime):
        result["updated_at"] = result["updated_at"].isoformat()

    src = result.get("source") or "WEBSITE"
    fid = result.get("form_id") or ""
    st = result.get("status") or "NEW"

    result["source_label"] = SOURCE_LABELS.get(src, src)
    result["form_label"] = FORM_LABELS.get(fid, fid.replace("_", " ").title())
    result["status_label"] = STATUS_CONFIG.get(st, {}).get("label", st)
    result["status_color"] = STATUS_CONFIG.get(st, {}).get("color", "#64748B")

    raw_msg = str(result.get("message") or result.get("note") or "").strip()
    structured: dict = dict(result.get("structured_details") or {})
    is_json_payload = False
    customer_actual_note = ""

    # Check if raw_msg is JSON (e.g. consignment submission)
    if (raw_msg.startswith("{") and raw_msg.endswith("}")) or raw_msg.startswith("{\n") or raw_msg.startswith("{\"") or raw_msg.startswith("{ \""):
        try:
            pj = json.loads(raw_msg)
            if isinstance(pj, dict):
                is_json_payload = True
                prop = pj.get("property") if isinstance(pj.get("property"), dict) else {}

                owner_role = pj.get("owner_role") or pj.get("role")
                if owner_role:
                    structured["Vai trò người gửi"] = str(owner_role)

                tx_type = prop.get("transaction_type") or pj.get("transaction_type")
                if tx_type:
                    structured["Hình thức giao dịch"] = format_transaction_type(tx_type)

                p_type = prop.get("property_type") or pj.get("property_type")
                if p_type:
                    structured["Loại bất động sản"] = format_property_type(p_type)

                exp_price = prop.get("expected_price") or pj.get("expected_price") or prop.get("price")
                if exp_price:
                    structured["Giá kỳ vọng"] = format_price_value(exp_price)

                exp_sale = prop.get("expected_sale_price") or pj.get("expected_sale_price")
                if exp_sale:
                    structured["Giá bán kỳ vọng"] = format_price_value(exp_sale)

                exp_rent = prop.get("expected_rent_price") or pj.get("expected_rent_price")
                if exp_rent:
                    structured["Giá thuê kỳ vọng"] = format_price_value(exp_rent)

                addr = prop.get("address") or pj.get("address")
                if addr:
                    structured["Địa chỉ BĐS"] = str(addr)

                city = prop.get("city") or pj.get("city")
                if city and city not in str(addr or ""):
                    structured["Khu vực"] = str(city)

                gross_area = prop.get("gross_area") or pj.get("gross_area")
                if gross_area and gross_area != 1:
                    structured["Diện tích"] = f"{gross_area} m²"

                legal = prop.get("legal_status") or pj.get("legal_status")
                if legal and legal != "Chưa cung cấp":
                    structured["Pháp lý"] = str(legal)

                title = prop.get("title") or pj.get("title")
                if title:
                    structured["Tiêu đề hồ sơ"] = str(title)

                notes_candidates = [
                    prop.get("notes"),
                    prop.get("description"),
                    pj.get("notes"),
                    pj.get("description"),
                    pj.get("message"),
                    pj.get("note"),
                ]
                for c in notes_candidates:
                    if c and isinstance(c, str) and not c.strip().startswith("{"):
                        customer_actual_note = c.strip()
                        break
        except Exception:
            is_json_payload = False

    if not is_json_payload:
        for line in raw_msg.split("\n"):
            line_clean = line.strip()
            if not line_clean:
                continue
            if ":" in line_clean:
                k, v = line_clean.split(":", 1)
                k_clean = k.strip().lstrip("-").strip()
                v_clean = v.strip()
                if k_clean and v_clean and not k_clean.startswith("{"):
                    structured[k_clean] = v_clean
            elif line_clean.startswith("[") and line_clean.endswith("]"):
                structured["Phân loại"] = line_clean[1:-1]
        customer_actual_note = raw_msg

    if is_json_payload:
        result["note"] = customer_actual_note
        result["message"] = customer_actual_note
        summary_parts = []
        p_type = structured.get("Loại bất động sản")
        p_addr = structured.get("Địa chỉ BĐS")
        p_price = structured.get("Giá kỳ vọng") or structured.get("Giá bán kỳ vọng") or structured.get("Giá thuê kỳ vọng")
        if p_type:
            summary_parts.append(p_type)
        if p_addr:
            summary_parts.append(p_addr)
        if p_price:
            summary_parts.append(f"Giá: {p_price}")
        result["consignment_summary"] = " • ".join(summary_parts) if summary_parts else None
    else:
        result["consignment_summary"] = None

    result["structured_details"] = structured

    media_ids = result.get("consignment_media_ids") or []
    if media_ids:
        if media_map:
            result["consignment_media_urls"] = [
                media_map.get(mid, f"/api/public/marketing-media/{mid}")
                for mid in media_ids
            ]
        else:
            result["consignment_media_urls"] = [
                f"/api/public/marketing-media/{mid}"
                for mid in media_ids
            ]
    else:
        result["consignment_media_urls"] = []

    return result


class WebsiteLeadService:
    @property
    def collection(self):
        return get_database()["website_leads"]

    async def create(self, payload, idempotency_key: str | None, request) -> dict:
        now = utcnow()
        database = get_database()
        if idempotency_key:
            fingerprint = hashlib.sha256(payload.model_dump_json().encode("utf-8")).hexdigest()
            try:
                await database["idempotency_records"].insert_one({
                    "key": f"lead:{idempotency_key}", "fingerprint": fingerprint,
                    "created_at": now, "expires_at": now + timedelta(hours=24),
                })
            except DuplicateKeyError:
                previous = await database["idempotency_records"].find_one({"key": f"lead:{idempotency_key}"})
                if previous and previous.get("fingerprint") == fingerprint and previous.get("reference_id"):
                    return {"reference_id": previous["reference_id"]}
                raise HTTPException(status_code=409, detail="Idempotency-Key đã được sử dụng")

        project_id = payload.interest_project_id or payload.project_id
        project = None
        if project_id and ObjectId.is_valid(project_id):
            project = await database["projects"].find_one({"_id": ObjectId(project_id), "is_published": True})
            if not project:
                project = await database["marketing_projects"].find_one({"_id": ObjectId(project_id), "is_published": True})

        # Resolve correct source
        src = payload.source
        if payload.form_id == "consignment" or payload.type == "CONSIGNMENT":
            src = "CONSIGNMENT"
        elif project_id:
            src = "PROJECT"

        lead_id = ObjectId()
        outbox_id = ObjectId()
        document = {
            "_id": lead_id,
            "version": 1,
            "customer_name": payload.customer_name,
            "phone": payload.phone,
            "normalized_phone": payload.phone,
            "email": payload.email.strip().lower() if payload.email else None,
            "message": payload.message(),
            "inquiry_type": payload.type,
            "source": src,
            "form_id": payload.form_id,
            "page_key": payload.page_key,
            "page_url": payload.page_url,
            "project_id": project_id,
            "project_name_snapshot": (project or {}).get("name"),
            "staff_profile_id": payload.staff_profile_id,
            "utm_source": payload.utm_source,
            "utm_medium": payload.utm_medium,
            "utm_campaign": payload.utm_campaign,
            "utm_content": payload.utm_content,
            "referrer": payload.referrer,
            "structured_details": payload.structured_details,
            "consent": payload.consent,
            "consent_version": payload.consent_version or "website-privacy-v1",
            "consented_at": now,
            "status": "NEW",
            "submitted_at": now,
            "created_at": now,
            "updated_at": now,
        }
        reference_id = str(lead_id)
        try:
            async with website_transaction() as session:
                kwargs = session_kwargs(session)
                await database["website_leads"].insert_one(document, **kwargs)
                await database["outbox_events"].insert_one({
                    "_id": outbox_id,
                    "event_type": "WEBSITE_LEAD_CREATED", "payload": {"lead_id": reference_id},
                    "status": "PENDING", "attempts": 0, "available_at": now,
                    "created_at": now, "updated_at": now,
                }, **kwargs)
                if idempotency_key:
                    await database["idempotency_records"].update_one(
                        {"key": f"lead:{idempotency_key}"},
                        {"$set": {"reference_id": reference_id, "status": "COMPLETED", "completed_at": now}},
                        **kwargs,
                    )
        except Exception:
            if not settings.MARKETING_TRANSACTIONS_ENABLED:
                await database["outbox_events"].delete_one({"_id": outbox_id})
                await database["website_leads"].delete_one({"_id": lead_id})
            if idempotency_key:
                # The reservation is intentionally created before the transaction so
                # concurrent requests cannot both enter it. Release an incomplete
                # reservation after either transactional rollback or local cleanup.
                await database["idempotency_records"].delete_one(
                    {"key": f"lead:{idempotency_key}", "reference_id": {"$exists": False}}
                )
            raise
        return {"reference_id": reference_id}

    async def rollback_create(
        self,
        reference_id: str,
        idempotency_key: str | None,
        media_ids: list[str] | None = None,
    ) -> None:
        """Compensate a multi-step public submission that did not complete."""
        database = get_database()
        if ObjectId.is_valid(reference_id):
            await database["website_leads"].delete_one({"_id": ObjectId(reference_id)})
        await database["outbox_events"].delete_many(
            {"event_type": "WEBSITE_LEAD_CREATED", "payload.lead_id": reference_id}
        )
        if idempotency_key:
            await database["idempotency_records"].delete_one(
                {"key": f"lead:{idempotency_key}", "reference_id": reference_id}
            )
        if media_ids:
            cleanup_time = utcnow()
            await database["media_assets"].update_many(
                {"_id": {"$in": media_ids}, "references": {"$size": 0}},
                {
                    "$set": {
                        "lifecycle_status": "ORPHANED",
                        "orphaned_at": cleanup_time,
                        "updated_at": cleanup_time,
                    }
                },
            )

    def query(
        self,
        *,
        status=None,
        source=None,
        form_id=None,
        project_id=None,
        search=None,
        date_from=None,
        date_to=None,
        ids=None,
    ):
        query: dict = {}
        if ids:
            query["_id"] = {"$in": [ObjectId(item) for item in ids if ObjectId.is_valid(item)]}
        if status:
            query["status"] = status
        if source:
            query["source"] = source
        if form_id:
            query["form_id"] = form_id
        if project_id:
            query["project_id"] = project_id
        if search:
            safe = re.escape(str(search).strip()[:100])
            query["$or"] = [
                {"customer_name": {"$regex": safe, "$options": "i"}},
                {"phone": {"$regex": safe}},
                {"email": {"$regex": safe, "$options": "i"}},
                {"message": {"$regex": safe, "$options": "i"}},
                {"project_name_snapshot": {"$regex": safe, "$options": "i"}},
            ]
        submitted: dict = {}
        if date_from:
            submitted["$gte"] = datetime.fromisoformat(str(date_from).replace("Z", "+00:00"))
        if date_to:
            submitted["$lte"] = datetime.fromisoformat(str(date_to).replace("Z", "+00:00"))
        if submitted:
            query["submitted_at"] = submitted
        return query

    async def list(self, query: dict, page: int, limit: int) -> dict:
        limit = max(1, min(limit, 200))
        page = max(1, page)
        total = await self.collection.count_documents(query)
        items = await self.collection.find(query).sort("submitted_at", -1).skip((page - 1) * limit).limit(limit).to_list(length=limit)

        all_media_ids = []
        for item in items:
            for mid in (item.get("consignment_media_ids") or []):
                all_media_ids.append(mid)

        media_map = {}
        if all_media_ids:
            cursor = get_database()["media_assets"].find({"_id": {"$in": all_media_ids}}, {"_id": 1, "url": 1})
            async for doc in cursor:
                file_id = str(doc["_id"])
                media_map[file_id] = doc.get("url") or f"/api/public/marketing-media/{file_id}"
            for mid in all_media_ids:
                if mid not in media_map:
                    media_map[mid] = f"/api/public/marketing-media/{mid}"

        return {
            "items": [serialize(item, media_map) for item in items],
            "total": total,
            "page": page,
            "limit": limit,
        }

    async def get_counts(self) -> dict:
        total = await self.collection.count_documents({})
        sources = ["COWORKING", "OFFICE", "PROJECT", "CONSIGNMENT", "WEBSITE"]
        by_source: dict[str, int] = {}
        for s in sources:
            by_source[s] = await self.collection.count_documents({"source": s})

        statuses = ["NEW", "VIEWED", "EXPORTED", "ARCHIVED"]
        by_status: dict[str, int] = {}
        for st in statuses:
            by_status[st] = await self.collection.count_documents({"status": st})

        return {
            "total": total,
            "by_source": by_source,
            "by_status": by_status,
        }

    async def update_status(self, lead_id: str, status: str, expected_version: int | None, actor: dict) -> dict:
        if not ObjectId.is_valid(lead_id):
            raise HTTPException(status_code=404, detail="Không tìm thấy thông tin khách hàng")
        existing = await self.collection.find_one({"_id": ObjectId(lead_id)})
        if not existing:
            raise HTTPException(status_code=404, detail="Không tìm thấy thông tin khách hàng")
        if existing.get("status") == status:
            return serialize(existing)

        query = {"_id": ObjectId(lead_id)}
        if expected_version is not None:
            query["version"] = expected_version

        doc = await self.collection.find_one_and_update(
            query,
            {"$set": {"status": status, "updated_at": utcnow(), "updated_by": actor.get("account_id")}, "$inc": {"version": 1}},
            return_document=ReturnDocument.AFTER,
        )
        if not doc:
            current = await self.collection.find_one({"_id": ObjectId(lead_id)})
            raise HTTPException(
                status_code=409,
                detail={
                    "code": "VERSION_CONFLICT",
                    "message": "Thông tin khách hàng đã được cập nhật ở phiên khác.",
                    "current_version": (current or {}).get("version"),
                },
            )
        return serialize(doc)

    async def export(self, query: dict, output_format: str, actor: dict) -> tuple[bytes, str, str]:
        items = await self.collection.find(query).sort("submitted_at", -1).limit(50000).to_list(length=50000)

        # Collect consignment media IDs to calculate resolved URLs
        all_media_ids = []
        for item in items:
            for mid in (item.get("consignment_media_ids") or []):
                all_media_ids.append(mid)

        media_map = {}
        if all_media_ids:
            cursor = get_database()["media_assets"].find({"_id": {"$in": all_media_ids}}, {"_id": 1, "url": 1})
            async for doc in cursor:
                file_id = str(doc["_id"])
                media_map[file_id] = doc.get("url") or f"/api/public/marketing-media/{file_id}"
            for mid in all_media_ids:
                if mid not in media_map:
                    media_map[mid] = f"/api/public/marketing-media/{mid}"

        headers = [
            "Mã yêu cầu (ID)",
            "Thời gian gửi",
            "Họ và tên khách hàng",
            "Số điện thoại",
            "Email",
            "Phân nhóm khách",
            "Biểu mẫu tiếp nhận",
            "Tóm tắt nhu cầu",
            "Loại BĐS / Dịch vụ",
            "Hình thức giao dịch",
            "Giá kỳ vọng / Ngân sách",
            "Địa chỉ BĐS / Tòa nhà",
            "Doanh nghiệp / Đơn vị",
            "Số chỗ / Diện tích",
            "Vai trò người gửi",
            "Số lượng ảnh BĐS",
            "Liên kết ảnh BĐS (URL xem trực tiếp)",
            "Dự án BĐS quan tâm",
            "Lời nhắn từ khách hàng",
            "Kênh tiếp thị (UTM)",
            "Chiến dịch (Campaign)",
            "Trang gửi biểu mẫu",
            "Trạng thái xử lý",
        ]

        serialized_items = [serialize(item, media_map) for item in items]
        rows = []

        for s in serialized_items:
            det = s.get("structured_details") or {}

            # Format submission datetime
            sub_at = s.get("submitted_at") or ""
            if isinstance(sub_at, str) and sub_at:
                try:
                    dt = datetime.fromisoformat(sub_at.replace("Z", "+00:00"))
                    sub_at_str = dt.strftime("%d/%m/%Y %H:%M:%S")
                except Exception:
                    sub_at_str = sub_at
            else:
                sub_at_str = str(sub_at)

            # Defensive customer message extraction (never allow raw JSON strings)
            raw_c_msg = str(s.get("message") or s.get("note") or "").strip()
            if (raw_c_msg.startswith("{") and raw_c_msg.endswith("}")) or raw_c_msg.startswith("{\n") or raw_c_msg.startswith("{\""):
                try:
                    parsed_c = json.loads(raw_c_msg)
                    prop_c = parsed_c.get("property") or {}
                    clean_c_msg = str(
                        prop_c.get("notes")
                        or prop_c.get("description")
                        or parsed_c.get("notes")
                        or parsed_c.get("description")
                        or ""
                    ).strip()
                except Exception:
                    clean_c_msg = ""
            else:
                clean_c_msg = raw_c_msg

            # Consignment & General details
            prop_type = (
                det.get("Loại bất động sản")
                or det.get("Gói thuê quan tâm")
                or det.get("Hình thức chỗ ngồi")
                or ""
            )
            tx_type = det.get("Hình thức giao dịch") or ""
            price_budget = (
                det.get("Giá kỳ vọng")
                or det.get("Giá bán kỳ vọng")
                or det.get("Giá thuê kỳ vọng")
                or det.get("Ngân sách dự kiến")
                or ""
            )
            addr_bldg = (
                det.get("Địa chỉ BĐS")
                or det.get("Tòa nhà quan tâm")
                or ""
            )
            company_name = det.get("Tên doanh nghiệp") or ""
            capacity_area = (
                det.get("Số lượng chỗ ngồi dự kiến")
                or det.get("Diện tích")
                or det.get("Diện tích / Mã sàn")
                or ""
            )
            if capacity_area and capacity_area.isdigit():
                capacity_area = f"{capacity_area} chỗ"

            sender_role = det.get("Vai trò người gửi") or ""
            media_urls = s.get("consignment_media_urls") or []
            media_count_str = f"{len(media_urls)} ảnh" if media_urls else "0"
            media_urls_str = "\n".join(media_urls) if media_urls else ""

            # Concise summary column
            summary_str = s.get("consignment_summary") or ""
            if not summary_str:
                parts = []
                if company_name:
                    parts.append(company_name)
                if prop_type:
                    parts.append(prop_type)
                if addr_bldg:
                    parts.append(addr_bldg)
                if price_budget:
                    parts.append(price_budget)
                summary_str = " • ".join(parts)

            rows.append([safe_spreadsheet_value(value) for value in [
                s.get("id") or "",
                sub_at_str or "",
                s.get("customer_name") or "",
                s.get("phone") or "",
                s.get("email") or "",
                s.get("source_label") or s.get("source") or "",
                s.get("form_label") or s.get("form_id") or "",
                summary_str or "",
                prop_type or "",
                tx_type or "",
                price_budget or "",
                addr_bldg or "",
                company_name or "",
                capacity_area or "",
                sender_role or "",
                media_count_str or "0",
                media_urls_str or "",
                s.get("project_name_snapshot") or "",
                clean_c_msg or "",
                s.get("utm_source") or "direct",
                s.get("utm_campaign") or "",
                s.get("page_url") or "",
                s.get("status_label") or s.get("status") or "",
            ]])

        if output_format == "csv":
            stream = io.StringIO()
            writer = csv.writer(stream)
            writer.writerow(headers)
            writer.writerows(rows)
            content = ("\ufeff" + stream.getvalue()).encode("utf-8")
            media_type, extension = "text/csv; charset=utf-8", "csv"
        else:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
            from openpyxl.utils import get_column_letter

            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "DS Khach Hang"

            # Header style
            header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="2D1B14", end_color="2D1B14", fill_type="solid")
            header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

            thin_border = Border(
                left=Side(style="thin", color="E2E8F0"),
                right=Side(style="thin", color="E2E8F0"),
                top=Side(style="thin", color="E2E8F0"),
                bottom=Side(style="thin", color="E2E8F0"),
            )

            sheet.row_dimensions[1].height = 28
            sheet.append(headers)

            for col_idx in range(1, len(headers) + 1):
                cell = sheet.cell(row=1, column=col_idx)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment

            # Data rows style
            row_font = Font(name="Segoe UI", size=10)
            for row_idx, row_data in enumerate(rows, start=2):
                sheet.row_dimensions[row_idx].height = 22
                sheet.append(row_data)
                for col_idx in range(1, len(row_data) + 1):
                    c = sheet.cell(row=row_idx, column=col_idx)
                    c.font = row_font
                    c.border = thin_border
                    if col_idx in (1, 2, 4, 10, 14, 16, 23):
                        c.alignment = Alignment(horizontal="center", vertical="center")
                    else:
                        c.alignment = Alignment(horizontal="left", vertical="center")

            # Auto-fit column widths
            for col in sheet.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = get_column_letter(col[0].column)
                sheet.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

            excel_binary = io.BytesIO()
            workbook.save(excel_binary)

            if output_format == "zip":
                zip_buffer = io.BytesIO()
                with zipfile.ZipFile(zip_buffer, "w", compression=zipfile.ZIP_DEFLATED) as zip_file:
                    zip_file.writestr("danh-sach-khach-hang.xlsx", excel_binary.getvalue())

                    for idx, s in enumerate(serialized_items, start=1):
                        m_ids = s.get("consignment_media_ids") or []
                        if not m_ids:
                            continue

                        c_name = re.sub(r'[^\w\s-]', '', s.get("customer_name") or f"Khach_{idx}").strip().replace(" ", "_")
                        c_phone = re.sub(r'[^\d+]', '', s.get("phone") or "NoPhone")
                        addr = re.sub(r'[^\w\s-]', '', (s.get("structured_details") or {}).get("Địa chỉ BĐS") or "").strip().replace(" ", "_")
                        folder_name = f"{c_phone}_{c_name}"
                        if addr:
                            folder_name += f"_{addr[:30]}"

                        for img_idx, mid in enumerate(m_ids, start=1):
                            try:
                                img_bytes, meta = await marketing_media_service.read(mid)
                                orig_name = meta.get("original_filename") or f"anh_{img_idx}.jpg"
                                clean_fname = re.sub(r'[^\w.-]', '_', orig_name)
                                arcname = f"Hinh_Anh_BDS_Ky_Gui/{folder_name}/anh_{img_idx}_{clean_fname}"
                                zip_file.writestr(arcname, img_bytes)
                            except Exception as e:
                                logger.warning(f"Could not include media {mid} in export zip: {e}")

                content = zip_buffer.getvalue()
                media_type, extension = "application/zip", "zip"
            else:
                content = excel_binary.getvalue()
                media_type, extension = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"

        now = utcnow()
        await get_database()["export_audits"].insert_one({
            "actor_id": actor.get("account_id") or actor.get("id"), "format": output_format,
            "row_count": len(items), "filters": query, "created_at": now,
        })
        if items:
            await self.collection.update_many(
                {"_id": {"$in": [item["_id"] for item in items]}},
                {
                    "$set": {
                        "status": "EXPORTED",
                        "exported_at": now,
                        "exported_by": actor.get("account_id"),
                        "updated_at": now,
                    },
                    "$inc": {"version": 1},
                },
            )
        prefix = "danh-sach-ky-gui-kem-anh" if output_format == "zip" else "danh-sach-khach-hang"
        return content, media_type, f"{prefix}_{now.date().isoformat()}.{extension}"

    async def download_lead_images(self, lead_id: str) -> tuple[bytes, str]:
        if not ObjectId.is_valid(lead_id):
            raise HTTPException(status_code=404, detail="Không tìm thấy thông tin khách hàng")
        lead = await self.collection.find_one({"_id": ObjectId(lead_id)})
        if not lead:
            raise HTTPException(status_code=404, detail="Không tìm thấy thông tin khách hàng")

        m_ids = lead.get("consignment_media_ids") or []
        if not m_ids:
            raise HTTPException(status_code=400, detail="Khách hàng này không có ảnh bất động sản đính kèm")

        c_name = re.sub(r'[^\w\s-]', '', lead.get("customer_name") or "KhachHang").strip().replace(" ", "_")
        c_phone = re.sub(r'[^\d+]', '', lead.get("phone") or "NoPhone")

        zip_buffer = io.BytesIO()
        with zipfile.ZipFile(zip_buffer, "w", compression=zipfile.ZIP_DEFLATED) as zip_file:
            for img_idx, mid in enumerate(m_ids, start=1):
                try:
                    img_bytes, meta = await marketing_media_service.read(mid)
                    orig_name = meta.get("original_filename") or f"anh_{img_idx}.jpg"
                    clean_fname = re.sub(r'[^\w.-]', '_', orig_name)
                    arcname = f"anh_{img_idx}_{clean_fname}"
                    zip_file.writestr(arcname, img_bytes)
                except Exception as e:
                    logger.warning(f"Could not read media {mid} for single lead zip: {e}")

        zip_content = zip_buffer.getvalue()
        filename = f"anh-ky-gui_{c_name}_{c_phone}.zip"
        return zip_content, filename


website_lead_service = WebsiteLeadService()
